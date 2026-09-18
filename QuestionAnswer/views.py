import os

from django.shortcuts import render
from django.conf import settings
from pypdf import PdfReader
from docx import Document as DocxDocument
from openai import OpenAI
from .models import PDFUpload, GeneratedResult


def extract_text_from_pdf(pdf_file):
    pdf_reader = PdfReader(pdf_file)
    extracted_text = ""
    for page in pdf_reader.pages:
        text = page.extract_text()
        if text:
            extracted_text += text + "\n"
    return extracted_text


def extract_text_from_docx(docx_file):
    document = DocxDocument(docx_file)
    extracted_text = ""

    # Paragraph text
    for paragraph in document.paragraphs:
        if paragraph.text:
            extracted_text += paragraph.text + "\n"

    # Text inside tables (often skipped, but notes/handouts use tables a lot)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    extracted_text += cell.text + "\n"

    return extracted_text


def extract_text_from_file(uploaded_file):
    """
    Dispatches to the right extractor based on file extension.
    Returns (text, error_message). error_message is None on success.
    """
    filename = uploaded_file.name.lower()
    ext = os.path.splitext(filename)[1]

    if ext == ".pdf":
        return extract_text_from_pdf(uploaded_file), None

    if ext == ".docx":
        return extract_text_from_docx(uploaded_file), None

    if ext == ".doc":
        # python-docx only reads the modern .docx (XML zip) format.
        # Legacy binary .doc files need conversion (e.g. LibreOffice/antiword)
        # before they can be parsed, so we reject them cleanly instead of
        # crashing on a bad parse.
        return None, "Legacy .doc files aren't supported yet — please save/export it as .docx and re-upload."

    return None, "Unsupported file type. Please upload a PDF or DOCX file."


def generate_questions(request):
    result_text = None
    error_message = None

    if request.method == "POST" and request.FILES.get("course_file"):
        uploaded_file = request.FILES["course_file"]
        num_mcqs = int(request.POST.get("num_mcqs", 5))
        num_questions = int(request.POST.get("num_questions", 5))

        # 1. Save upload record in DB
        pdf_record = PDFUpload.objects.create(
            file=uploaded_file,
            num_mcqs=num_mcqs,
            num_questions=num_questions,
        )

        # 2. Extract text from the uploaded file (PDF or DOCX)
        file_text, extract_error = extract_text_from_file(uploaded_file)

        if extract_error:
            error_message = extract_error
        elif not file_text or not file_text.strip():
            error_message = "Unable to extract text from the file. Please try another file."
        else:
            try:
                # 3. Build prompt
                prompt = f"""
                You are an expert educator. Based on the provided text extracted from a document, generate two separate sections of questions.

                SECTION 1: MULTIPLE CHOICE QUESTIONS (MCQs)
                Generate exactly {num_mcqs} MCQs. For each MCQ, provide:
                - The question text
                - Four distinct options (A, B, C, D)
                - The correct answer with a brief one-line explanation.

                SECTION 2: SUBJECTIVE QUESTIONS (SHORT/LONG ANSWER)
                Generate exactly {num_questions} conceptual/subjective questions that test understanding of the text. For each question, provide:
                - The question text
                - A brief model answer or ideal response guideline based on the text.

                Format the entire output cleanly using proper Markdown headers and bullet points.

                Source Text:
                {file_text[:5000]}
                """

                # 4. Call local OpenAI-compatible server
                client = OpenAI(
                    api_key=settings.MODEL_API_KEY,
                    base_url=settings.MODEL_API_BASE_URL,
                    timeout=settings.MODEL_API_TIMEOUT,
                )

                response = client.chat.completions.create(
                    model=settings.MODEL_NAME,
                    messages=[
                        {"role": "system", "content": "You are an expert educator who writes clear exam questions."},
                        {"role": "user", "content": prompt}
                    ],
                    temperature=0.7,
                )

                result_text = response.choices[0].message.content

                # 5. Save result in DB
                GeneratedResult.objects.create(pdf=pdf_record, result_text=result_text)

            except Exception as e:
                error_message = f"An error occurred: {e}"

    return render(request, "question.html", {
        "result_text": result_text,
        "error_message": error_message,
    })